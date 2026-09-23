import sqlite3
from pathlib import Path

import openpyxl
import pandas as pd

from src.screener.engine import run_preset
from src.analytics.peer import calculate_percentile


DB_PATH = Path("nifty100.db")
SCREENER_PATH = Path("output/screener_output.xlsx")
PEER_PATH = Path("output/peer_comparison.xlsx")
RADAR_DIR = Path("reports/radar_charts")

def test_screener_output_exists():
    assert SCREENER_PATH.exists()


def test_peer_comparison_output_exists():
    assert PEER_PATH.exists()


def test_radar_chart_count():
    charts = list(RADAR_DIR.glob("*_radar.png"))
    assert len(charts) == 92


def test_screener_has_six_sheets():
    workbook = openpyxl.load_workbook(
        SCREENER_PATH,
        read_only=True,
    )

    expected = {
        "Quality Compounder",
        "Value Pick",
        "Growth Accelerator",
        "Dividend Champion",
        "Debt-Free Blue Chip",
        "Turnaround Watch",
    }

    assert set(workbook.sheetnames) == expected


def test_peer_comparison_has_eleven_sheets():
    workbook = openpyxl.load_workbook(
        PEER_PATH,
        read_only=True,
    )

    assert len(workbook.sheetnames) == 11


def test_peer_percentiles_has_560_rows():
    conn = sqlite3.connect(DB_PATH)

    try:
        count = conn.execute(
            "SELECT COUNT(*) FROM peer_percentiles"
        ).fetchone()[0]
    finally:
        conn.close()

    assert count == 560


def test_peer_percentiles_has_ten_metrics():
    conn = sqlite3.connect(DB_PATH)

    try:
        count = conn.execute(
            "SELECT COUNT(DISTINCT metric) FROM peer_percentiles"
        ).fetchone()[0]
    finally:
        conn.close()

    assert count == 10
def test_it_services_highest_roe_has_highest_percentile():
    conn = sqlite3.connect(DB_PATH)

    try:
        query = """
            SELECT
                company_id,
                value,
                percentile_rank
            FROM peer_percentiles
            WHERE peer_group_name = 'IT Services'
              AND metric = 'ROE'
        """

        df = pd.read_sql_query(query, conn)

    finally:
        conn.close()

    df = df.dropna(subset=["value", "percentile_rank"])

    highest_roe = df.loc[
        df["value"].idxmax()
    ]

    highest_percentile = df.loc[
        df["percentile_rank"].idxmax()
    ]

    assert highest_roe["company_id"] == highest_percentile["company_id"]


def test_it_services_de_inversion():
    conn = sqlite3.connect(DB_PATH)

    try:
        query = """
            SELECT
                company_id,
                value,
                percentile_rank
            FROM peer_percentiles
            WHERE peer_group_name = 'IT Services'
              AND metric = 'D/E'
        """

        df = pd.read_sql_query(query, conn)

    finally:
        conn.close()

    df = df.dropna(
        subset=["value", "percentile_rank"]
    )

    lowest_de = df.loc[
        df["value"].idxmin()
    ]

    highest_percentile = df.loc[
        df["percentile_rank"].idxmax()
    ]

    assert lowest_de["company_id"] == highest_percentile["company_id"]


def test_fmcg_highest_roe_has_highest_percentile():
    conn = sqlite3.connect(DB_PATH)

    try:
        query = """
            SELECT
                company_id,
                value,
                percentile_rank
            FROM peer_percentiles
            WHERE peer_group_name = 'FMCG'
              AND metric = 'ROE'
        """

        df = pd.read_sql_query(query, conn)

    finally:
        conn.close()

    df = df.dropna(
        subset=["value", "percentile_rank"]
    )

    highest_roe = df.loc[
        df["value"].idxmax()
    ]

    highest_percentile = df.loc[
        df["percentile_rank"].idxmax()
    ]

    assert highest_roe["company_id"] == highest_percentile["company_id"]


def test_percentiles_are_between_zero_and_hundred():
    conn = sqlite3.connect(DB_PATH)

    try:
        df = pd.read_sql_query(
            """
            SELECT percentile_rank
            FROM peer_percentiles
            WHERE percentile_rank IS NOT NULL
            """,
            conn,
        )

    finally:
        conn.close()

    assert df["percentile_rank"].between(
        0,
        100,
    ).all()


def test_peer_comparison_has_22_columns():
    workbook = openpyxl.load_workbook(
        PEER_PATH,
        read_only=True,
    )

    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        assert ws.max_column == 22


def test_peer_comparison_has_benchmark_and_summary():
    workbook = openpyxl.load_workbook(
        PEER_PATH,
        read_only=True,
    )

    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]

        values = []

        for row in ws.iter_rows(
            values_only=True
        ):
            values.extend(
                value
                for value in row
                if value is not None
            )

        assert "BENCHMARK" in values
        assert "SUMMARY" in values


def test_quality_compounder_thresholds():
    result = run_preset(
        "quality_compounder"
    )

    assert not result.empty

    assert (
        result["return_on_equity_pct"] > 15
    ).all()

    assert (
        result["debt_to_equity"] < 1
    ).all()