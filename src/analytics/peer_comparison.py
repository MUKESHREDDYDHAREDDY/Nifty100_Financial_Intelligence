from pathlib import Path
import sqlite3

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter


DB_PATH = Path("nifty100.db")
OUTPUT_PATH = Path("output/peer_comparison.xlsx")
PEER_METRICS = {
    "ROE": ("return_on_equity_pct", True),
    "ROCE": ("roce_percentage", True),
    "NPM": ("net_profit_margin_pct", True),
    "D/E": ("debt_to_equity", False),
    "FCF": ("free_cash_flow_cr", True),
    "PAT CAGR 5yr": ("pat_cagr_5yr", True),
    "Revenue CAGR 5yr": ("revenue_cagr_5yr", True),
    "EPS CAGR 5yr": ("eps_cagr_5yr", True),
    "Interest Coverage": ("interest_coverage", True),
    "Asset Turnover": ("asset_turnover", True),
}
def load_peer_comparison_data(db_path=DB_PATH):
    """
    Load latest annual financial data, company names,
    peer-group assignments, and benchmark information.
    """

    conn = sqlite3.connect(db_path)

    try:
        query = """
            SELECT
                f.company_id,
                c.company_name,
                f.year,
                c.roce_percentage,
                f.return_on_equity_pct,
                f.net_profit_margin_pct,
                f.debt_to_equity,
                f.free_cash_flow_cr,
                f.pat_cagr_5yr,
                f.revenue_cagr_5yr,
                f.eps_cagr_5yr,
                f.interest_coverage,
                f.asset_turnover,
                f.composite_quality_score,
                pg.peer_group_name,
                pg.is_benchmark
            FROM financial_ratios f
            JOIN companies c
                ON f.company_id = c.id
            LEFT JOIN peer_groups pg
                ON f.company_id = pg.company_id
        """

        df = pd.read_sql_query(query, conn)

    finally:
        conn.close()

    # Exclude TTM records.
    df = df[df["year"] != "TTM"].copy()

    # Keep the latest annual record for each company.
    df = (
        df
        .sort_values("year")
        .groupby("company_id", sort=False)
        .tail(1)
        .reset_index(drop=True)
    )

    return df

def load_percentile_data(db_path=DB_PATH):
    """
    Load the peer percentile table created on Day 18.
    """

    conn = sqlite3.connect(db_path)

    try:
        query = """
            SELECT
                company_id,
                peer_group_name,
                metric,
                value,
                percentile_rank,
                year
            FROM peer_percentiles
        """

        df = pd.read_sql_query(query, conn)

    finally:
        conn.close()

    return df

def build_peer_comparison_data(
    financial_data,
    percentile_data,
):
    """
    Combine financial values with their peer percentile ranks.
    """

    value_columns = [
        "company_id",
        "company_name",
        "year",
        "peer_group_name",
        "is_benchmark",
    ]

    for metric_name, (column_name, _) in PEER_METRICS.items():
        if column_name not in value_columns:
            value_columns.append(column_name)

    values = financial_data[value_columns].copy()

    # Convert the long percentile table into one row per company.
    percentile_wide = percentile_data.pivot_table(
        index=[
            "company_id",
            "peer_group_name",
        ],
        columns="metric",
        values="percentile_rank",
        aggfunc="first",
    ).reset_index()

    percentile_wide.columns.name = None

    percentile_columns = {
        metric: "{} Percentile".format(metric)
        for metric in PEER_METRICS
        if metric in percentile_wide.columns
    }

    percentile_wide = percentile_wide.rename(
        columns=percentile_columns
    )

    result = values.merge(
        percentile_wide,
        on=[
            "company_id",
            "peer_group_name",
        ],
        how="left",
    )

    return result

def apply_percentile_color(cell):
    """
    Apply percentile-based color coding.

    Green  = >= 75
    Yellow = > 25 and < 75
    Red    = <= 25
    """

    value = cell.value

    if value is None:
        return

    try:
        value = float(value)
    except (TypeError, ValueError):
        return

    if value >= 75:
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="C6EFCE",
        )

    elif value <= 25:
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="FFC7CE",
        )

    else:
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="FFEB9C",
        )
def format_peer_sheet(
    ws,
    percentile_columns,
):
    """
    Format a peer-group worksheet.
    """

    # Header formatting
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # Freeze header row
    ws.freeze_panes = "A2"

    # Apply percentile colors
    for column_index in percentile_columns:
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(
                row=row,
                column=column_index,
            )

            apply_percentile_color(cell)

    # Adjust column widths
    for column_cells in ws.columns:
        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = 0

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value)),
                )

        ws.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 12),
            30,
        )

    # Enable filtering
    ws.auto_filter.ref = ws.dimensions
def export_peer_comparison(
    db_path=DB_PATH,
    output_path=OUTPUT_PATH,
):
    """
    Create peer_comparison.xlsx with exactly 11 peer-group sheets.
    """

    financial_data = load_peer_comparison_data(db_path)
    percentile_data = load_percentile_data(db_path)

    data = build_peer_comparison_data(
        financial_data,
        percentile_data,
    )

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()

    # Remove the default worksheet.
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    created_sheets = []

    for peer_group_name in sorted(
        data["peer_group_name"]
        .dropna()
        .unique()
    ):

        group = data[
            data["peer_group_name"] == peer_group_name
        ].copy()

        # Create a valid Excel sheet name.
        sheet_name = str(peer_group_name)[:31]

        ws = workbook.create_sheet(
            title=sheet_name
        )

        created_sheets.append(sheet_name)

        # ------------------------------------------
        # Build columns
        # ------------------------------------------

        columns = [
            "company_id",
            "company_name",
        ]

        percentile_columns = []

        for metric_name, (
            column_name,
            _,
        ) in PEER_METRICS.items():

            columns.append(column_name)

            percentile_name = "{} Percentile".format(
                metric_name
            )

            columns.append(percentile_name)

        # Write header
        for column_index, column_name in enumerate(
            columns,
            start=1,
        ):
            ws.cell(
                row=1,
                column=column_index,
                value=column_name,
            )

        # ------------------------------------------
        # Write company rows
        # ------------------------------------------

        row_number = 2

        for _, company in group.iterrows():

            for column_index, column_name in enumerate(
                columns,
                start=1,
            ):

                value = company.get(
                    column_name
                )

                ws.cell(
                    row=row_number,
                    column=column_index,
                    value=value,
                )

                if "Percentile" in column_name:
                    percentile_columns.append(
                        column_index
                    )

            row_number += 1

        # ------------------------------------------
        # Benchmark row
        # ------------------------------------------

        benchmark = group[
            group["is_benchmark"] == 1
        ]

        if not benchmark.empty:

            benchmark_row = row_number

            ws.cell(
                row=benchmark_row,
                column=1,
                value="BENCHMARK",
            )

            ws.cell(
                row=benchmark_row,
                column=2,
                value="Benchmark",
            )

            for column_index, column_name in enumerate(
                columns[2:],
                start=3,
            ):

                if column_name in benchmark.columns:

                    value = benchmark.iloc[0][
                        column_name
                    ]

                    ws.cell(
                        row=benchmark_row,
                        column=column_index,
                        value=value,
                    )

            # Amber/gold benchmark formatting
            for cell in ws[benchmark_row]:
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="FFD966",
                )
                cell.font = Font(
                    bold=True
                )

            row_number += 1

        # ------------------------------------------
        # Peer-group median summary row
        # ------------------------------------------

        summary_row = row_number

        ws.cell(
            row=summary_row,
            column=1,
            value="SUMMARY",
        )

        ws.cell(
            row=summary_row,
            column=2,
            value="Peer Group Median",
        )

        for column_index, column_name in enumerate(
            columns[2:],
            start=3,
        ):

            if column_name in group.columns:

                numeric_values = pd.to_numeric(
                    group[column_name],
                    errors="coerce",
                )

                median_value = numeric_values.median()

                ws.cell(
                    row=summary_row,
                    column=column_index,
                    value=median_value,
                )

        for cell in ws[summary_row]:
            cell.font = Font(
                bold=True
            )

        # ------------------------------------------
        # Format worksheet
        # ------------------------------------------

        format_peer_sheet(
            ws,
            sorted(set(percentile_columns)),
        )

    workbook.save(output_path)

    return {
        "output_path": str(output_path),
        "sheet_count": len(created_sheets),
        "sheets": created_sheets,
    }
if __name__ == "__main__":
    result = export_peer_comparison()

    print("Output:", result["output_path"])
    print("Sheet count:", result["sheet_count"])
    print("Sheets:", result["sheets"])